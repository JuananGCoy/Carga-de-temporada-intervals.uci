import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_excel_template():
    filename = "Plantilla_Entrenamientos.xlsx"
    
    # 1. Definir los datos de ejemplo
    data = {
        "Día": [
            "2026-06-08",
            "2026-06-09",
            "2026-06-10",
            "2026-06-11",
            "2026-06-12"
        ],
        "Deporte": [
            "Bici",
            "Carrera",
            "Natación",
            "Gimnasio",
            "Carrera"
        ],
        "Nombre del Entreno": [
            "Series de FTP (Ejemplo)",
            "Rodaje Z2 (Ejemplo)",
            "Técnica y Umbral (Ejemplo)",
            "Rutina Fuerza A (Ejemplo)",
            "Test 10k (Ejemplo)"
        ],
        "Entrenamiento principal": [
            "Warmup\n- 20m Z1\n\nMain Set\n3x\n  - 8m Z4\n  - 4m Z1\n\nCooldown\n- 10m Z1",
            "Warmup\n- 10m Z1 Pace\n\nMain Set\n- 40m Z2 Pace\n\nCooldown\n- 5m Z1 Pace",
            "Warmup\n- 400m Z1\n\nMain Set\n- 10x 100m Z3 r: 20s\n\nCooldown\n- 200m Z1",
            "- 50m\n\nBloque Fuerza:\n- 4x8 Sentadillas\n- 4x8 Peso Muerto\n- 3x12 Zancadas\n- Core: 3 rondas de Rueda Abdominal",
            "TEST 10K - Ritmo umbral.\nHacer calentamiento libre previo y luego correr 10 kilómetros a ritmo máximo constante."
        ],
        "Categoría": [
            "WORKOUT",
            "WORKOUT",
            "WORKOUT",
            "WORKOUT",
            "RACE_A"
        ]
    }
    
    # Crear DataFrame
    df = pd.DataFrame(data)
    
    # Guardar en excel
    df.to_excel(filename, sheet_name="Entrenos", index=False)
    
    # 2. Estilizar el Excel con openpyxl para que se vea profesional y premium
    wb = openpyxl.load_workbook(filename)
    ws = wb["Entrenos"]
    
    # Fuentes y Estilos
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="333333")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Azul oscuro
    data_fill_zebra = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid") # Gris muy claro para filas pares
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    # Configurar Cabecera
    for col_idx, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # Configurar Datos
    for row_idx in range(2, len(data["Día"]) + 2):
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
            # Cebra en filas pares
            if row_idx % 2 == 0:
                cell.fill = data_fill_zebra
                
            # Alinear fecha y deporte al centro
            if col_idx in [1, 2, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
    # Ajustar anchos de columnas
    ws.column_dimensions['A'].width = 15  # Día
    ws.column_dimensions['B'].width = 15  # Deporte
    ws.column_dimensions['C'].width = 30  # Nombre del Entreno
    ws.column_dimensions['D'].width = 50  # Entrenamiento principal (texto largo)
    ws.column_dimensions['E'].width = 15  # Categoría
    
    # Ajustar altura de fila de datos
    for r in range(2, len(data["Día"]) + 2):
        ws.row_dimensions[r].height = 100 # Espacio para el texto con saltos de línea
        
    ws.row_dimensions[1].height = 28 # Altura cabecera
    
    # Guardar los cambios
    wb.save(filename)
    print(f"Plantilla creada exitosamente: '{filename}'")

if __name__ == "__main__":
    create_excel_template()
